from pathlib import Path
import re
import pickle

import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from pypdf import PdfReader
from docx import Document

INDEX_FILE = Path(__file__).parent / "data" / "knowledge_index.pkl"


def extract_pdf(path):
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = re.sub(r"\s+", " ", text).strip()
        if text:
            pages.append(f"[Page {i}] {text}")
    return "\n".join(pages)


def extract_docx(path):
    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        t = re.sub(r"\s+", " ", p.text).strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [re.sub(r"\s+", " ", c.text).strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(path):
    if path.suffix.lower() == ".pdf":
        return extract_pdf(path)
    if path.suffix.lower() == ".docx":
        return extract_docx(path)
    return ""


def chunk_text(text, chunk_size=1200, overlap=200):
    words = text.split()
    chunks = []
    start = 0
    while start < len(words):
        end = min(len(words), start + chunk_size)
        chunks.append(" ".join(words[start:end]))
        if end == len(words):
            break
        start = end - overlap
    return chunks


def load_documents(folder):
    return sorted(
        p.name for p in Path(folder).glob("*")
        if p.suffix.lower() in {".pdf", ".docx"}
    )


def build_index(folder):
    folder = Path(folder)
    records = []

    for path in sorted(folder.glob("*")):
        if path.suffix.lower() not in {".pdf", ".docx"}:
            continue
        text = extract_text(path)
        if not text.strip():
            # Scanned/image-only files are retained in the folder but not indexed.
            continue
        for n, chunk in enumerate(chunk_text(text)):
            records.append(
                {"source": path.name, "chunk": n + 1, "text": chunk}
            )

    if not records:
        raise RuntimeError(
            "No extractable text was found. Scanned PDFs require OCR before indexing."
        )

    vectorizer = TfidfVectorizer(
        lowercase=True,
        stop_words="english",
        ngram_range=(1, 2),
        max_features=50000,
    )
    matrix = vectorizer.fit_transform([r["text"] for r in records])

    index = {
        "records": records,
        "vectorizer": vectorizer,
        "matrix": matrix,
    }

    INDEX_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(INDEX_FILE, "wb") as f:
        pickle.dump(index, f)

    return index


def search_index(index, query, k=6):
    q = index["vectorizer"].transform([query])
    scores = cosine_similarity(q, index["matrix"]).ravel()
    order = np.argsort(scores)[::-1][:k]

    results = []
    for i in order:
        if scores[i] <= 0:
            continue
        r = index["records"][i].copy()
        r["score"] = float(scores[i])
        results.append({
            "source": f"{r['source']} — chunk {r['chunk']}",
            "text": r["text"],
            "score": r["score"],
        })
    return results
